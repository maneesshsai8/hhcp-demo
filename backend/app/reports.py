"""
Scorecard exports — Excel (OpenPyXL) and PDF (Playwright HTML->PDF, with a
pure-Python fpdf2 fallback if Chromium isn't available). This is the demo's
stand-in for the blueprint's "document/spreadsheet generation" subsystem.
"""
import io
import html as _html

# Status values that should read as "good" (green) vs "bad" (red) in exports.
_GOOD = {"on_track", "complete", "solved", "done"}
_BAD = {"off_track", "open"}


def _fmt(v):
    if v is None:
        return ""
    return str(v)


# --------------------------------------------------- generic list exports ----
# Used by Rocks / Issues / To-Dos. `columns` is a list of (header, key); a
# column keyed "status" gets green/red cell shading.

def build_table_xlsx(sheet_title: str, headline: str, columns: list[tuple], rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    navy = "1B3A5C"
    head_fill = PatternFill("solid", fgColor=navy)
    head_font = Font(color="FFFFFF", bold=True)
    good_fill = PatternFill("solid", fgColor="E7F3EC")
    bad_fill = PatternFill("solid", fgColor="FBEAE7")
    thin = Side(style="thin", color="DFE2DC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = headline
    ws["A1"].font = Font(size=15, bold=True, color=navy)

    row0 = 3
    for c, (header, _key) in enumerate(columns, start=1):
        cell = ws.cell(row=row0, column=c, value=header)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="left")
        cell.border = border

    for i, r in enumerate(rows):
        rr = row0 + 1 + i
        for c, (_header, key) in enumerate(columns, start=1):
            cell = ws.cell(row=rr, column=c, value=_fmt(r.get(key)))
            cell.border = border
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=(key == "description"))
            if key == "status":
                v = (r.get("status") or "").lower()
                if v in _GOOD:
                    cell.fill = good_fill
                elif v in _BAD:
                    cell.fill = bad_fill

    widths = {"title": 34, "description": 40, "owner_name": 20, "created_by_name": 20, "team_name": 18}
    for c, (_header, key) in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(row=row0, column=c).column_letter].width = widths.get(key, 14)
    ws.freeze_panes = "A4"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_table_html(headline: str, columns: list[tuple], rows: list[dict]) -> str:
    head = "".join(f"<th>{_html.escape(h)}</th>" for h, _ in columns)

    def cell(r, key):
        val = _html.escape(_fmt(r.get(key)))
        if key == "status":
            v = (r.get("status") or "").lower()
            cls = "on" if v in _GOOD else ("off" if v in _BAD else "")
            return f'<td class="{cls}">{val}</td>'
        return f'<td class="{"kpi" if key == "title" else ""}">{val}</td>'

    body = "".join("<tr>" + "".join(cell(r, key) for _, key in columns) + "</tr>" for r in rows)
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>
      body {{ font-family: -apple-system, 'Segoe UI', sans-serif; color:#1b2430; padding:28px; }}
      h1 {{ color:#122943; font-size:20px; margin:0 0 14px; }}
      table {{ border-collapse:collapse; width:100%; font-size:11px; }}
      th, td {{ border:1px solid #dfe2dc; padding:6px 8px; text-align:left; vertical-align:top; }}
      thead th {{ background:#1b3a5c; color:#fff; }}
      td.kpi {{ font-weight:600; color:#122943; }}
      td.on {{ background:#e7f3ec; }} td.off {{ background:#fbeae7; }}
    </style></head><body>
      <h1>{_html.escape(headline)}</h1>
      <table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>
    </body></html>"""


async def build_table_pdf(headline: str, columns: list[tuple], rows: list[dict]) -> tuple[bytes, str]:
    html = build_table_html(headline, columns, rows)
    try:
        from playwright.async_api import async_playwright
        async with async_playwright() as p:
            browser = await p.chromium.launch()
            page = await browser.new_page()
            await page.set_content(html, wait_until="load")
            pdf = await page.pdf(format="A4", landscape=True, print_background=True,
                                 margin={"top": "12mm", "bottom": "12mm", "left": "10mm", "right": "10mm"})
            await browser.close()
        return pdf, "playwright"
    except Exception:
        return _build_table_pdf_fallback(headline, columns, rows), "fpdf2"


def _build_table_pdf_fallback(headline: str, columns: list[tuple], rows: list[dict]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.add_page()
    pdf.set_text_color(18, 41, 67)
    pdf.set_font("Helvetica", "B", 15)
    pdf.cell(0, 10, headline.encode("latin-1", "replace").decode("latin-1"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    avail = 277
    col_w = avail / len(columns)
    pdf.set_font("Helvetica", "B", 8)
    pdf.set_fill_color(27, 58, 92)
    pdf.set_text_color(255, 255, 255)
    for header, _key in columns:
        pdf.cell(col_w, 7, header, border=1, fill=True)
    pdf.ln()

    pdf.set_font("Helvetica", "", 8)
    for r in rows:
        for _header, key in columns:
            txt = _fmt(r.get(key))[:40].encode("latin-1", "replace").decode("latin-1")
            if key == "status":
                v = (r.get("status") or "").lower()
                pdf.set_fill_color(*(231, 243, 236) if v in _GOOD else ((251, 234, 231) if v in _BAD else (255, 255, 255)))
                pdf.set_text_color(18, 41, 67)
                pdf.cell(col_w, 6, txt, border=1, fill=True)
            else:
                pdf.set_fill_color(255, 255, 255)
                pdf.set_text_color(18, 41, 67)
                pdf.cell(col_w, 6, txt, border=1)
        pdf.ln()

    return bytes(pdf.output())


# ------------------------------------------------------------------ Excel ----
def build_scorecard_xlsx(tenant_name: str, scorecards: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Scorecard"

    navy = "1B3A5C"
    on_fill = PatternFill("solid", fgColor="E7F3EC")
    off_fill = PatternFill("solid", fgColor="FBEAE7")
    mid_fill = PatternFill("solid", fgColor="FCF3D9")
    head_fill = PatternFill("solid", fgColor=navy)
    head_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="DFE2DC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = f"Scorecard — {tenant_name}"
    ws["A1"].font = Font(size=15, bold=True, color=navy)

    # widest week history across KPIs -> column set
    max_weeks = max((len(s["weekly_history"]) for s in scorecards), default=0)
    weeks = []
    for s in scorecards:
        for w in s["weekly_history"]:
            if w["week_ending"] not in weeks:
                weeks.append(w["week_ending"])
    weeks = sorted(weeks)[-max_weeks:] if max_weeks else []

    header = ["KPI", "Owner", "Target"] + weeks
    row0 = 3
    for c, title in enumerate(header, start=1):
        cell = ws.cell(row=row0, column=c, value=title)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    for i, s in enumerate(scorecards):
        r = row0 + 1 + i
        ws.cell(row=r, column=1, value=s["title"]).border = border
        ws.cell(row=r, column=2, value=s["owner"] or "Unassigned").border = border
        ws.cell(row=r, column=3, value=f'{s["comparison_operator"]} {s["target_value"]} {s["unit"] or ""}'.strip()).border = border
        by_week = {w["week_ending"]: w for w in s["weekly_history"]}
        for c, wk in enumerate(weeks, start=4):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.alignment = Alignment(horizontal="center")
            w = by_week.get(wk)
            if w:
                cell.value = w["actual_value"]
                rag = w.get("rag") or ("GREEN" if w["status"] == "ON_TRACK" else "RED")
                cell.fill = {"GREEN": on_fill, "YELLOW": mid_fill, "RED": off_fill}.get(rag, off_fill)

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 16
    for c in range(4, 4 + len(weeks)):
        ws.column_dimensions[ws.cell(row=row0, column=c).column_letter].width = 12
    ws.freeze_panes = "D4"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ HTML -----
def build_scorecard_html(tenant_name: str, scorecards: list[dict]) -> str:
    weeks = sorted({w["week_ending"] for s in scorecards for w in s["weekly_history"]})

    def cells(s):
        by_week = {w["week_ending"]: w for w in s["weekly_history"]}
        out = ""
        for wk in weeks:
            w = by_week.get(wk)
            if w:
                rag = w.get("rag") or ("GREEN" if w["status"] == "ON_TRACK" else "RED")
                cls = {"GREEN": "on", "YELLOW": "mid", "RED": "off"}.get(rag, "off")
                out += f'<td class="{cls}">{w["actual_value"]}</td>'
            else:
                out += "<td></td>"
        return out

    rows = "".join(
        f"<tr><td class='kpi'>{_html.escape(s['title'])}</td>"
        f"<td>{_html.escape(s['owner'] or 'Unassigned')}</td>"
        f"<td>{_html.escape(s['comparison_operator'])} {s['target_value']} {_html.escape(s['unit'] or '')}</td>"
        f"{cells(s)}</tr>"
        for s in scorecards
    )
    week_cols = "".join(f"<th>{wk[5:]}</th>" for wk in weeks)  # MM-DD
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>
      body {{ font-family: -apple-system, 'Segoe UI', sans-serif; color:#1b2430; padding:28px; }}
      h1 {{ color:#122943; font-size:20px; margin:0 0 2px; }}
      .sub {{ color:#5b6570; font-size:12px; margin:0 0 18px; }}
      table {{ border-collapse:collapse; width:100%; font-size:11px; }}
      th, td {{ border:1px solid #dfe2dc; padding:5px 7px; text-align:center; }}
      thead th {{ background:#1b3a5c; color:#fff; }}
      td.kpi {{ text-align:left; font-weight:600; color:#122943; }}
      td.on {{ background:#e7f3ec; }} td.mid {{ background:#fcf3d9; }} td.off {{ background:#fbeae7; }}
    </style></head><body>
      <h1>Scorecard — {_html.escape(tenant_name)}</h1>
      <p class="sub">Measurables with Red / Yellow / Green status.</p>
      <table><thead><tr><th style="text-align:left">KPI</th><th>Owner</th><th>Target</th>{week_cols}</tr></thead>
      <tbody>{rows}</tbody></table>
    </body></html>"""


# ------------------------------------------------------------------ PDF ------
async def build_scorecard_pdf(tenant_name: str, scorecards: list[dict]) -> tuple[bytes, str]:
    """
    Returns (pdf_bytes, engine). Tries Playwright (headless Chromium rendering
    the HTML template — the blueprint's approach); falls back to fpdf2 (pure
    Python) so the endpoint works even without a browser installed.
    """
    html = build_scorecard_html(tenant_name, scorecards)
    try:
        from playwright.async_api import async_playwright
        async with async_playwright() as p:
            browser = await p.chromium.launch()
            page = await browser.new_page()
            await page.set_content(html, wait_until="load")
            pdf = await page.pdf(format="A4", landscape=True, print_background=True,
                                 margin={"top": "12mm", "bottom": "12mm", "left": "10mm", "right": "10mm"})
            await browser.close()
        return pdf, "playwright"
    except Exception:
        return _build_scorecard_pdf_fallback(tenant_name, scorecards), "fpdf2"


def _build_scorecard_pdf_fallback(tenant_name: str, scorecards: list[dict]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.add_page()
    pdf.set_text_color(18, 41, 67)
    pdf.set_font("Helvetica", "B", 15)
    pdf.cell(0, 10, f"Scorecard - {tenant_name}", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(91, 101, 112)
    pdf.cell(0, 6, "Weekly measurables (on-track / off-track)", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    weeks = sorted({w["week_ending"] for s in scorecards for w in s["weekly_history"]})
    kpi_w, tgt_w = 55, 28
    wk_w = max(12, (277 - kpi_w - tgt_w) / max(len(weeks), 1))

    pdf.set_font("Helvetica", "B", 8)
    pdf.set_fill_color(27, 58, 92)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(kpi_w, 7, "KPI", border=1, fill=True)
    pdf.cell(tgt_w, 7, "Target", border=1, align="C", fill=True)
    for wk in weeks:
        pdf.cell(wk_w, 7, wk[5:], border=1, align="C", fill=True)
    pdf.ln()

    pdf.set_font("Helvetica", "", 8)
    for s in scorecards:
        by_week = {w["week_ending"]: w for w in s["weekly_history"]}
        pdf.set_text_color(18, 41, 67)
        pdf.set_fill_color(255, 255, 255)
        pdf.cell(kpi_w, 6, s["title"][:34], border=1)
        pdf.cell(tgt_w, 6, f'{s["comparison_operator"]} {s["target_value"]}', border=1, align="C")
        for wk in weeks:
            w = by_week.get(wk)
            if w:
                rag = w.get("rag") or ("GREEN" if w["status"] == "ON_TRACK" else "RED")
                rgb = {"GREEN": (231, 243, 236), "YELLOW": (252, 243, 217), "RED": (251, 234, 231)}.get(rag, (251, 234, 231))
                pdf.set_fill_color(*rgb)
                pdf.cell(wk_w, 6, str(w["actual_value"]), border=1, align="C", fill=True)
            else:
                pdf.set_fill_color(255, 255, 255)
                pdf.cell(wk_w, 6, "", border=1, fill=True)
        pdf.ln()

    return bytes(pdf.output())
